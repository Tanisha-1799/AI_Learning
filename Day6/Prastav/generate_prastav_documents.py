"""Generate the 5 required fictional Prastav source documents."""
from pathlib import Path
import json
import csv

import pymupdf as fitz
import docx

DOCS_DIR = Path("documents")
DOCS_DIR.mkdir(exist_ok=True)


def write_pdf_capability_statement() -> None:
    file_path = DOCS_DIR / "01_network_reliability_sla_capability_statement.pdf"

    page_text = """
Network Reliability & SLA Capability Statement
Document ID: NCS-CAP-001 | Version: 2.3 | Effective Date: 15 July 2026

NCS Telco+ has delivered a measured network uptime of 99.99% across the last 36 months
(2023-2025 audited service reports) for managed enterprise network estates.

Mean Time to Repair (MTTR) performance has been sustained at 41 minutes for P1 incidents,
measured over the same period across all managed contracts.

Certification and compliance profile:
- ISO/IEC 27001:2022 certified security management program
- SOC 2 Type II controls attestation for managed operations workflows
- Annual business continuity drills aligned to ISO 22301 practices

For RFP submissions, these figures may be cited as historical delivery evidence and should
be accompanied by relevant service scope qualifiers.
""".strip()

    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    rect = fitz.Rect(50, 50, 545, 792)
    page.insert_textbox(rect, page_text, fontsize=11, fontname="helv", align=0)
    pdf.save(file_path)
    pdf.close()


def write_docx_case_study() -> None:
    file_path = DOCS_DIR / "02_client_case_study_orion_retail.docx"

    d = docx.Document()
    d.add_heading("Client Case Study - Orion Retail Group", level=1)
    d.add_paragraph("Document ID: NCS-CASE-002 | Version: 1.7 | Effective Date: 20 June 2026")

    d.add_heading("Client Starting Problem", level=2)
    d.add_paragraph(
        "Orion Retail operated 140 stores with frequent WAN instability, fragmented monitoring, "
        "and inconsistent escalation ownership. Average outage restoration exceeded 90 minutes."
    )

    d.add_heading("Solution Delivered", level=2)
    d.add_paragraph(
        "NCS Telco+ implemented managed SD-WAN with centralized observability, automated "
        "incident triage, and a unified service management runbook with named L2/L3 ownership."
    )

    d.add_heading("Quantified Results", level=2)
    d.add_paragraph("- Reduced critical network downtime by 35% within the first two quarters.")
    d.add_paragraph("- Cut support tickets by 40% through proactive fault detection and remediation.")
    d.add_paragraph("- Improved MTTR from 96 minutes to 54 minutes for priority incidents.")

    d.add_heading("Why This Matters for RFP Responses", level=2)
    d.add_paragraph(
        "This case demonstrates measurable modernization outcomes and supports claims tied to "
        "service continuity, operational maturity, and support efficiency."
    )

    d.save(file_path)


def write_html_sales_faq() -> None:
    file_path = DOCS_DIR / "03_rfp_response_playbook_sales_faq.html"
    html = """<!doctype html>
<html lang=\"en\">
<head>
  <meta charset=\"utf-8\" />
  <title>RFP Response Playbook / Sales FAQ</title>
</head>
<body>
  <h1>RFP Response Playbook / Sales FAQ</h1>
  <p>Document ID: NCS-RFP-FAQ-003 | Version: 1.4 | Effective Date: 10 August 2026</p>

  <h2>Who approves non-standard pricing in an RFP?</h2>
  <p>
    Non-standard pricing is approved by the Commercial Approval Board (CAB) chaired by
    the Regional Sales Director. For strategic bids above SGD 1M, final sign-off is by
    the Vice President, Enterprise Sales.
  </p>

  <h2>What is the standard turnaround time for technical RFP clarifications?</h2>
  <p>
    Standard turnaround is 2 business days for routine technical clarifications and
    4 business days for architecture-level clarifications requiring platform SMEs.
  </p>

  <h2>Sales Engineering contact for urgent bid support</h2>
  <p>
    Contact: Priya Menon, Lead Sales Engineer
    Email: priya.menon@ncs-telcoplus.example
    Phone: +65 9123 4567
  </p>
</body>
</html>
"""
    file_path.write_text(html, encoding="utf-8")


def write_csv_win_loss() -> None:
    file_path = DOCS_DIR / "04_rfp_win_loss_log.csv"
    fieldnames = [
        "rfp_id",
        "client_name",
        "submission_date",
        "outcome",
        "primary_reason",
        "deal_value_tier",
    ]
    rows = [
        ["RFP-2026-001", "Northshore Bank", "2026-02-18", "Won", "Superior migration timeline and phased cutover plan", "Tier-High"],
        ["RFP-2026-002", "MetroCare Health", "2026-03-05", "Lost", "Price uncompetitive against incumbent", "Tier-High"],
        ["RFP-2026-003", "Aster Logistics", "2026-03-29", "Won", "Strong managed service SLA commitments", "Tier-Medium"],
        ["RFP-2026-004", "Skyline Retail Group", "2026-04-14", "Lost", "Security compliance appendix submitted late", "Tier-Medium"],
        ["RFP-2026-005", "BluePeak Manufacturing", "2026-05-06", "Won", "Demonstrated multi-site rollout capability", "Tier-Low"],
        ["RFP-2026-006", "Horizon Hospitality", "2026-05-28", "Lost", "Insufficient local support coverage in proposal", "Tier-Low"],
        ["RFP-2026-007", "Vertex Education", "2026-06-19", "Won", "Best score on service continuity and governance model", "Tier-Medium"],
        ["RFP-2026-008", "Civic Transit Authority", "2026-07-11", "Lost", "Weak response detail on disaster recovery testing cadence", "Tier-High"],
    ]
    with file_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(fieldnames)
        writer.writerows(rows)


def write_json_service_catalogue() -> None:
    file_path = DOCS_DIR / "05_service_catalogue.json"
    payload = {
        "document_id": "NCS-SVC-CAT-005",
        "version": "2.1",
        "effective_date": "01 September 2026",
        "services": [
            {
                "service_name": "Managed Network Services",
                "category": "Operations",
                "description": "24x7 monitoring, incident response, and proactive network performance optimization.",
                "pricing_tier": "Premium",
                "sla_commitment": "99.95% monthly uptime with 30-minute P1 response",
            },
            {
                "service_name": "Enterprise SD-WAN",
                "category": "Connectivity",
                "description": "Policy-driven branch networking with centralized orchestration and secure overlays.",
                "pricing_tier": "Standard",
                "sla_commitment": "99.90% monthly uptime with 1-hour P1 response",
            },
            {
                "service_name": "Unified Communications Stack",
                "category": "Collaboration",
                "description": "Cloud calling, messaging, and meeting workloads integrated with identity controls.",
                "pricing_tier": "Standard",
                "sla_commitment": "99.90% monthly uptime with 1-hour major incident response",
            },
            {
                "service_name": "Zero Trust Access Gateway",
                "category": "Security",
                "description": "Context-aware access controls for users, devices, and applications.",
                "pricing_tier": "Premium",
                "sla_commitment": "99.95% monthly uptime with 30-minute critical incident response",
            },
        ],
    }
    file_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def main() -> None:
    write_pdf_capability_statement()
    write_docx_case_study()
    write_html_sales_faq()
    write_csv_win_loss()
    write_json_service_catalogue()
    print("Generated Prastav documents in ./documents")


if __name__ == "__main__":
    main()
